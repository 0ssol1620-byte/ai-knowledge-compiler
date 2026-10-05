"""Merged XLSX source extents survive CIR without exceeding parser fences."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from akc_cir import CanonicalDocument, canonical_json
from akc_native_parsers import (
    ParseContext,
    ParserLimits,
    StructuredParseError,
    parse_non_pdf_to_cir,
    xlsx_parser,
)
from openpyxl import Workbook


def _payload(values: dict[str, str | int], merges: tuple[str, ...]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Table 9"
    for coordinate, value in values.items():
        sheet[coordinate] = value
    for reference in merges:
        sheet.merge_cells(reference)
    output = io.BytesIO()
    book.save(output)
    book.close()
    assert len(output.getvalue()) < 8192
    return output.getvalue()


def _parse(payload: bytes, limits: ParserLimits | None = None) -> CanonicalDocument:
    document = parse_non_pdf_to_cir(
        filename="merge.xlsx",
        declared_mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        data=payload,
        limits=limits,
        context=ParseContext(
            tenant_id="tenant_k13",
            document_id="document_k13",
            document_version_id="version_k13",
            created_at=datetime(2026, 10, 5, tzinfo=UTC),
        ),
    )
    wire = canonical_json(document)
    restored = CanonicalDocument.model_validate_json(wire)
    assert restored == document and canonical_json(restored) == wire
    return restored


@pytest.mark.parametrize(
    ("values", "merges", "bounds", "shape", "anchor", "position", "span", "count"),
    [
        (
            {"A1": "Title", "Q22": 0, "A21": "Note"},
            ("A21:R21",),
            "A1:R22",
            (22, 18),
            "A21",
            (20, 0),
            (1, 18),
            3,
        ),
        ({"C4": "Header", "C5": "Note"}, ("C5:H7",), "C4:H7", (4, 6), "C5", (1, 0), (3, 6), 2),
        ({"C4": "Note"}, ("C4:F4",), "C4:F4", (1, 4), "C4", (0, 0), (1, 4), 1),
        ({"C4": "Note"}, ("C4:C8",), "C4:C8", (5, 1), "C4", (0, 0), (5, 1), 1),
        ({}, ("C4:F8",), "C4:F8", (5, 4), "C4", (0, 0), (5, 4), 1),
        (
            {"A1": "Title", "D4": "Note", "B7": "Other", "Q22": 0},
            ("D4:F5", "B7:C8"),
            "A1:Q22",
            (22, 17),
            "D4",
            (3, 3),
            (2, 3),
            4,
        ),
    ],
)
def test_merge_extents_preserve_source_anchors_without_invented_cells(
    values,
    merges,
    bounds,
    shape,
    anchor,
    position,
    span,
    count,
) -> None:
    document = _parse(_payload(values, merges))
    (table,) = [block.table for block in document.blocks if block.table]
    assert (table.row_count, table.column_count) == shape
    assert table.source_refs[0].native_object_id == f"xlsx/sheet/0000/range/{bounds}"
    assert len(table.cells) == count
    cells = {cell.source_refs[0].native_object_id.rsplit("/", 1)[-1]: cell for cell in table.cells}
    assert set(cells) == set(values) | {reference.split(":")[0] for reference in merges}
    cell = cells[anchor]
    assert (cell.row_index0, cell.column_index0) == position
    assert (cell.row_span, cell.column_span) == span
    assert cell.raw_text == str(values.get(anchor, ""))
    if anchor not in values:
        assert cell.value_type is None
    if "Q22" in cells:
        assert cells["Q22"].raw_text == "0" and cells["Q22"].value_type == "number"


def _rewrite_merges(references: tuple[str, ...]) -> bytes:
    payload = _payload({"B2": "Note"}, ("B2:C3",))
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(payload)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for entry in source.infolist():
            data = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                merges = "".join(f'<mergeCell ref="{reference}"/>' for reference in references)
                data = re.sub(
                    rb"<mergeCells\b.*?</mergeCells>",
                    f'<mergeCells count="{len(references)}">{merges}</mergeCells>'.encode(),
                    data,
                )
            target.writestr(entry, data)
    return output.getvalue()


@pytest.mark.parametrize(
    ("references", "changes", "code"),
    [
        (("A1:A100001",), {}, "SHEET_ROW_LIMIT"),
        (("A1:XFD1",), {}, "SHEET_COLUMN_LIMIT"),
        (("A1:Z20000",), {}, "SHEET_CELL_LIMIT"),
        (("A1:H8",), {"max_table_cells": 63}, "TABLE_CELL_LIMIT"),
        (("B2:B6",), {"max_table_rows": 4}, "TABLE_ROW_LIMIT"),
        (("B2:F2",), {"max_table_columns": 4}, "TABLE_COLUMN_LIMIT"),
        (("B2:C3", "F2:G3"), {"max_table_cells": 7}, "TABLE_CELL_LIMIT"),
        (("A0:B1",), {}, "XLSX_INVALID_RANGE"),
        (("C3:A1",), {}, "XLSX_INVALID_RANGE"),
    ],
)
def test_hostile_merge_rejected_before_openpyxl_expansion(monkeypatch, references, changes, code):
    payload = _rewrite_merges(references)

    def forbidden_load(*args, **kwargs):
        pytest.fail("merge must be rejected before openpyxl expands it")

    monkeypatch.setattr(xlsx_parser, "load_workbook", forbidden_load)
    with pytest.raises(StructuredParseError) as error:
        _parse(payload, replace(ParserLimits(), **changes))
    assert error.value.code == code


@pytest.mark.parametrize(
    ("changes", "code"),
    [
        ({"max_table_cells": 24}, "TABLE_CELL_LIMIT"),
        ({"max_table_rows": 4}, "TABLE_ROW_LIMIT"),
        ({"max_table_columns": 4}, "TABLE_COLUMN_LIMIT"),
    ],
)
@pytest.mark.parametrize("guard", ["openpyxl", "covered"])
def test_union_grid_budgets_checked_before_expansion(monkeypatch, changes, code, guard):
    payload = _payload({"A1": "Title", "D4": "Note"}, ("D4:E5",))

    def forbidden_expansion(*args, **kwargs):
        pytest.fail("union grid must be bounded before expansion")

    monkeypatch.setattr(
        xlsx_parser,
        "load_workbook" if guard == "openpyxl" else "range",
        forbidden_expansion,
        raising=False,
    )
    with pytest.raises(StructuredParseError) as error:
        _parse(payload, replace(ParserLimits(), **changes))
    assert error.value.code == code


def test_exact_union_grid_budget_accepts_preserved_offset_merge():
    limits = replace(ParserLimits(), max_table_rows=4, max_table_columns=6, max_table_cells=24)
    document = _parse(_payload({"C4": "Header", "C5": "Note"}, ("C5:H7",)), limits)
    table = next(block.table for block in document.blocks if block.table)
    assert (table.row_count, table.column_count, len(table.cells)) == (4, 6, 2)


@pytest.mark.parametrize("references", [("B2:C3", "C3:D4"), ("B2:C3", "B2:D4")])
def test_overlapping_merge_declarations_are_rejected_without_clipping_anchor(references):
    with pytest.raises(StructuredParseError) as error:
        _parse(_rewrite_merges(references))
    assert error.value.code == "TABLE_CELL_OVERLAP"
