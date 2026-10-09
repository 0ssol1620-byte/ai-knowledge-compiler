"""Tiny synthetic K13 cases through source validation, native parsing and CIR wire format."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from akc_cir import CanonicalCell, CanonicalDocument, CanonicalTable, canonical_json, sha256_digest
from akc_native_parsers import (
    ParseContext,
    ParserLimits,
    StructuredParseError,
    parse_non_pdf_to_cir,
)
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _parse(
    payload: bytes, extension: str = "xlsx", *, limits: ParserLimits | None = None
) -> CanonicalDocument:
    document = parse_non_pdf_to_cir(
        filename=f"k13.{extension}",
        declared_mime="text/csv" if extension == "csv" else _XLSX_MIME,
        data=payload,
        limits=limits,
        context=ParseContext(
            tenant_id="tenant_k13",
            document_id="document_k13",
            document_version_id="version_k13",
            created_at=datetime(2026, 10, 5, tzinfo=UTC),
        ),
    )
    assert document.source_sha256 == sha256_digest(payload)
    wire = canonical_json(document)
    restored = CanonicalDocument.model_validate_json(wire)
    assert restored == document
    assert canonical_json(restored) == wire
    return restored


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    payload = output.getvalue()
    assert len(payload) < 16_384
    return payload


def _tables(document: CanonicalDocument) -> list[CanonicalTable]:
    return [block.table for block in document.blocks if block.table is not None]


def _cells(table: CanonicalTable) -> dict[str | None, CanonicalCell]:
    return {cell.source_refs[0].native_object_id: cell for cell in table.cells}


def test_csv_duplicate_headers_multiline_unicode_and_blank_zero_keep_grid() -> None:
    # These are source strings, including decomposed Unicode and a quoted CRLF.
    payload = 'value,value,note\r\n0,,"  cafe\u0301\r\n둘째 줄  "\r\n"0",,=1+2\r\n'.encode()
    document = _parse(payload, "csv")
    (table,) = _tables(document)
    assert (table.row_count, table.column_count, table.header_row_count) == (3, 3, 0)
    assert table.source_refs[0].native_object_id == "csv/table/range/A1:C3"
    assert "csv_header_not_declared" in table.quality_flags
    cells = _cells(table)
    assert cells["csv/row/000000/cell/A1"].raw_text == "value"
    assert cells["csv/row/000000/cell/B1"].raw_text == "value"
    assert "csv/row/000001/cell/B2" not in cells
    assert "csv/row/000002/cell/B3" not in cells
    for row in (1, 2):
        zero = cells[f"csv/row/{row:06d}/cell/A{row + 1}"]
        assert (zero.row_index0, zero.column_index0, zero.raw_text, zero.value_type) == (
            row,
            0,
            "0",
            "string",
        )
    note = cells["csv/row/000001/cell/C2"]
    assert (note.raw_text, note.normalized_text) == ("  cafe\u0301\r\n둘째 줄  ", "café\n둘째 줄")
    formula_like = cells["csv/row/000002/cell/C3"]
    assert (formula_like.raw_text, formula_like.value_type, formula_like.formula) == (
        "=1+2",
        "string",
        None,
    )
    assert "spreadsheet_formula_prefix_preserved_as_text" in formula_like.quality_flags
    assert "csv_formula_prefixed_text_not_executed" in document.metadata["warnings"]
    assert all(cell.source_refs[0].bbox1000 is None for cell in table.cells)


def test_xlsx_offset_sheets_named_tables_duplicate_headers_dates_units_blank_zero() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "매출 e\u0301"
    for coordinate, value in {
        "C4": "Value",
        "D4": "Value",
        "E4": "Date",
        "F4": "Mass",
        "C5": 0,
        "D5": "",
        "E5": datetime(2024, 12, 31),
        "F5": 12.5,
        "C6": "0",
    }.items():
        sheet[coordinate] = value
    sheet["E5"].number_format = "yyyy-mm-dd"
    sheet["F5"].number_format = '0.0" kg"'
    declared = workbook.create_sheet("Declared")
    for coordinate, value in {
        "C4": "Region",
        "D4": "Value",
        "C5": "East",
        "D5": 2,
        "G4": "Region",
        "H4": "Value",
        "G5": "West",
        "H5": 3,
    }.items():
        declared[coordinate] = value
    declared.add_table(Table(displayName="EastResults", ref="C4:D5"))
    declared.add_table(Table(displayName="WestResults", ref="G4:H5"))
    document = _parse(_workbook_bytes(workbook))
    assert [item["name"] for item in document.metadata["sheets"]] == ["매출 e\u0301", "Declared"]
    first, second = _tables(document)
    assert first.id != second.id
    assert first.source_refs[0].native_object_id == "xlsx/sheet/0000/range/C4:F6"
    assert (first.row_count, first.column_count, first.header_row_count) == (3, 4, 1)
    assert "header_row_inferred" in first.quality_flags
    cells = _cells(first)
    assert cells["xlsx/sheet/0000/cell/C4"].raw_text == "Value"
    assert cells["xlsx/sheet/0000/cell/C4"].raw_text_verbatim is None
    assert cells["xlsx/sheet/0000/cell/D4"].raw_text == "Value"
    assert "xlsx/sheet/0000/cell/D5" not in cells
    zero = cells["xlsx/sheet/0000/cell/C5"]
    assert (zero.row_index0, zero.column_index0, zero.raw_text, zero.value_type) == (
        1,
        0,
        "0",
        "number",
    )
    assert zero.raw_text_verbatim is None
    assert cells["xlsx/sheet/0000/cell/C6"].value_type == "string"
    date = cells["xlsx/sheet/0000/cell/E5"]
    assert (date.raw_text, date.value_type, date.number_format) == (
        "2024-12-31T00:00:00",
        "datetime",
        "yyyy-mm-dd",
    )
    mass = cells["xlsx/sheet/0000/cell/F5"]
    assert (mass.raw_text, mass.value_type, mass.number_format) == ("12.5", "number", '0.0" kg"')
    assert second.source_refs[0].native_object_id == "xlsx/sheet/0001/range/C4:H5"
    assert document.metadata["sheets"][1]["tables"] == [
        {"name": "EastResults", "displayName": "EastResults", "ref": "C4:D5"},
        {"name": "WestResults", "displayName": "WestResults", "ref": "G4:H5"},
    ]
    for table in (first, second):
        for cell in table.cells:
            ref = cell.source_refs[0]
            assert (ref.document_id, ref.document_version_id) == ("document_k13", "version_k13")
            assert ref.page_number1 == ref.page_index0 + 1
            assert ref.bbox1000 is None


@pytest.mark.parametrize(
    ("source", "normalized"),
    [
        ("  padded  ", "padded"),
        ("cafe\u0301", "café"),
        ("first\r\nsecond", "first\nsecond"),
        (" \t ", ""),
    ],
)
def test_xlsx_raw_string_cell_survives_normalized_display_and_wire(
    source: str, normalized: str
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["C4"] = "Label"
    sheet["C5"] = source
    document = _parse(_workbook_bytes(workbook))
    (table,) = _tables(document)
    cell = _cells(table)["xlsx/sheet/0000/cell/C5"]
    assert cell.raw_text == source
    assert cell.normalized_text == normalized
    assert cell.raw_text_verbatim is True
    assert "xlsx_cell_text_normalized" in cell.quality_flags
    assert (cell.row_index0, cell.column_index0, cell.value_type) == (1, 0, "string")


@pytest.mark.parametrize(
    "values",
    [
        (" " * 1_000,),
        (" \r\n" * 500,),
        ("cafe\u0301" * 30,),
        (" " * 70, "\t" * 70),
        ("x" * 1_000,),
    ],
    ids=("spaces", "whitespace-crlf", "decomposed-unicode", "cumulative-cells", "nonblank"),
)
def test_xlsx_preserved_raw_text_cannot_bypass_total_text_limit(values: tuple[str, ...]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "S"
    sheet["A1"] = "H"
    for row, value in enumerate(values, start=2):
        sheet.cell(row, 1, value)
    with pytest.raises(StructuredParseError) as failure:
        _parse(
            _workbook_bytes(workbook),
            limits=replace(ParserLimits(), max_total_text_chars=128),
        )
    assert failure.value.code == "EXTRACTED_TEXT_LIMIT"


@pytest.mark.parametrize("raw_length", (125, 126, 127))
def test_xlsx_raw_text_budget_boundary_counts_sheet_and_display_text(raw_length: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "S"
    sheet["A1"] = "H"
    sheet["A2"] = " " * raw_length
    payload = _workbook_bytes(workbook)
    limits = replace(ParserLimits(), max_total_text_chars=128)
    # Existing output accounts for "S" + "H"; newly preserved source text
    # accounts for each space, cumulatively under the same total fence.
    if raw_length + 2 > limits.max_total_text_chars:
        with pytest.raises(StructuredParseError) as failure:
            _parse(payload, limits=limits)
        assert failure.value.code == "EXTRACTED_TEXT_LIMIT"
    else:
        document = _parse(payload, limits=limits)
        (table,) = _tables(document)
        cell = _cells(table)["xlsx/sheet/0000/cell/A2"]
        assert (cell.raw_text, cell.normalized_text) == (" " * raw_length, "")


@pytest.mark.parametrize(("value", "budget"), (("abcd", 5), ("=1+2", 9)))
def test_xlsx_existing_plain_and_formula_budget_boundaries_stay_unchanged(
    value: str, budget: int
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "S"
    sheet["A1"] = value
    payload = _workbook_bytes(workbook)
    document = _parse(payload, limits=replace(ParserLimits(), max_total_text_chars=budget))
    (table,) = _tables(document)
    cell = table.cells[0]
    assert (cell.raw_text, cell.normalized_text) == (value, value)
    assert cell.raw_text_verbatim is (True if value.startswith("=") else None)
    assert "xlsx_cell_text_normalized" not in cell.quality_flags
    with pytest.raises(StructuredParseError) as failure:
        _parse(payload, limits=replace(ParserLimits(), max_total_text_chars=budget - 1))
    assert failure.value.code == "EXTRACTED_TEXT_LIMIT"


def test_xlsx_zero_boolean_string_error_caches_and_uncached_formula_are_distinct() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for coordinate in ("C4", "C5", "C6", "C7", "C8"):
        sheet[coordinate] = "=1/0"
    payload = _workbook_bytes(workbook)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as archive, zipfile.ZipFile(output, "w") as target:
        for entry in archive.infolist():
            data = archive.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                for coordinate, value_type, cache in (
                    ("C4", "n", "0"),
                    ("C5", "b", "0"),
                    ("C6", "str", "cached text"),
                    ("C7", "e", "#DIV/0!"),
                ):
                    needle = f'<c r="{coordinate}"><f>1/0</f><v></v></c>'.encode()
                    assert data.count(needle) == 1
                    replacement = (
                        f'<c r="{coordinate}" t="{value_type}"><f>1/0</f><v>{cache}</v></c>'
                    ).encode()
                    data = data.replace(needle, replacement)
            target.writestr(entry, data)
    document = _parse(output.getvalue())
    (table,) = _tables(document)
    cells = _cells(table)
    for coordinate, display, value_type in (
        ("C4", "0", "number"),
        ("C5", "FALSE", "boolean"),
        ("C6", "cached text", "string"),
        ("C7", "#DIV/0!", "error"),
        ("C8", "=1/0", None),
    ):
        cell = cells[f"xlsx/sheet/0000/cell/{coordinate}"]
        assert (cell.raw_text, cell.formula) == ("=1/0", "=1/0")
        assert (cell.normalized_text, cell.value_type) == (display, value_type)
        assert "formula_preserved_not_executed" in cell.quality_flags
        assert ("formula_cached_value_missing" in cell.quality_flags) == (coordinate == "C8")
    formulas = {item["cell"]: item for item in document.metadata["sheets"][0]["formulas"]}
    assert formulas["C4"]["cachedValuePresent"] is True
    assert formulas["C4"]["cachedValue"] == "0"
    assert formulas["C8"]["cachedValuePresent"] is False
    assert formulas["C8"]["cachedValue"] is None


@pytest.mark.parametrize(
    ("member", "contents", "code"),
    [
        ("xl/vbaProject.bin", b"synthetic active-content marker", "OFFICE_ACTIVE_CONTENT"),
        (
            "xl/_rels/k13.xml.rels",
            (
                f'<Relationships xmlns="{_NS}"><Relationship Id="external" '
                'Type="http://schemas.openxmlformats.org/officeDocument/'
                '2006/relationships/externalLink" '
                'Target="https://example.invalid/k13.xlsx" TargetMode="External"/></Relationships>'
            ).encode(),
            "OFFICE_EXTERNAL_RELATION",
        ),
    ],
)
def test_xlsx_active_content_is_rejected_before_workbook_loading(
    member: str, contents: bytes, code: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    from akc_native_parsers import xlsx_parser

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "safe value"
    payload = _workbook_bytes(workbook)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as archive, zipfile.ZipFile(output, "w") as target:
        for entry in archive.infolist():
            target.writestr(entry, archive.read(entry))
        target.writestr(member, contents)

    def forbidden_load(*args: object, **kwargs: object) -> None:
        pytest.fail("untrusted active content reached workbook loading")

    monkeypatch.setattr(xlsx_parser, "load_workbook", forbidden_load)
    with pytest.raises(StructuredParseError) as failure:
        _parse(output.getvalue())
    assert failure.value.code == code
